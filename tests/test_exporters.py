# -*- coding: utf-8 -*-
"""tests/test_exporters.py — p2-3 导出中心：SRT 时间轴 / PDF / Word（离线秒级）。

覆盖：
- SRT 时间轴两档来源（precise sidecar vs derived 墙钟差分）、同秒退化段的钳制、
  时间戳格式化；纯英文笔记单行退化；
- sidecar 读写往返 + 损坏/版本不符容错；
- PDF（fpdf2 + msyh.ttc）与 DOCX（python-docx，含 w:eastAsia）真生成、产物头合法；
- asr.transcribe 的 with_times 三形态返回；
- GET /api/notes/export?fmt= 端到端（md/srt/pdf/docx + 非法 fmt 400 + 缺笔记 404
  + 导出失败给可见错误）。

PDF/DOCX 用真库真生成（不 mock）——依赖已在 requirements.txt 的 P2 段。
"""
import asyncio
import io
import json
from contextlib import asynccontextmanager
from pathlib import Path

import httpx
import numpy as np
import pytest
from types import SimpleNamespace

import exporters
import main
import notes as notes_mod

NOTE = "lecture_20260101_0900.md"


def write_note(name=NOTE, pairs=(("01:00:00", "Hello world.", "你好世界。"),
                                 ("01:00:05", "Second line.", "第二行。")),
               course="计算机科学导论"):
    main.TRANSCRIPTS.mkdir(parents=True, exist_ok=True)
    lines = ["# 双语课堂笔记 20260101_0900", "",
             "> LectureLens 自动生成（EN 原文 / ZH 翻译）"]
    if course:
        lines.append(f"> 课程：{course}")
    lines += ["", "---", ""]
    for i, (ts, en, zh) in enumerate(pairs, start=1):
        lines.append(f"- `{ts}` {en}")
        if zh:
            lines.append(f"  - `#{i}` {zh}")
    p = main.TRANSCRIPTS / name
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return p


def pairs_of(name=NOTE):
    return notes_mod.parse((main.TRANSCRIPTS / name).read_text(encoding="utf-8"))["pairs"]


@asynccontextmanager
async def asgi_client():
    transport = httpx.ASGITransport(app=main.app, client=("127.0.0.1", 123))
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


# —— SRT 基础函数 ——

def test_srt_timestamp_format():
    assert exporters.srt_timestamp(0) == "00:00:00,000"
    assert exporters.srt_timestamp(1.5) == "00:00:01,500"
    assert exporters.srt_timestamp(3661.25) == "01:01:01,250"
    assert exporters.srt_timestamp(-5) == "00:00:00,000"   # 负值按 0（时钟回拨兜底）


def test_clock_seconds_parsing():
    assert exporters.clock_seconds("01:46:42") == 1 * 3600 + 46 * 60 + 42
    assert exporters.clock_seconds("00:00:00") == 0
    for bad in ("", "1:2", "xx:yy:zz", "25:00:00", "00:61:00", "00:00:61", None, 12):
        assert exporters.clock_seconds(bad) is None, f"应拒绝 {bad!r}"


# —— 时间轴：derived（墙钟差分） ——

def test_times_for_derived_from_clock():
    write_note()
    pairs = pairs_of(NOTE)
    cues, source = exporters.times_for(pairs, None)
    assert source == "derived"
    assert cues[0][0] == 0.0                 # 原点 = 首段
    assert cues[0][1] == 5.0                 # 次段开始 = 本段结束
    assert cues[1][0] == 5.0
    assert cues[1][1] >= 5.0 + exporters.MIN_CUE_S


def test_times_for_derived_clamps_same_second_segments():
    """秒级分辨率下相邻段可能同秒（差 0）——裸用会出 0 时长 cue，必须钳制。"""
    main.TRANSCRIPTS.mkdir(parents=True, exist_ok=True)
    p = write_note(pairs=(("01:00:00", "a", "甲"), ("01:00:00", "b", "乙"),
                          ("01:00:00", "c", "丙")))
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    cues, _ = exporters.times_for(pairs, None)
    for st, en in cues:
        assert en - st >= exporters.MIN_CUE_S - 1e-9
    starts = [c[0] for c in cues]
    assert starts == sorted(starts) and len(set(starts)) == 3   # 严格递增


def test_times_for_caps_long_cue():
    p = write_note(pairs=(("01:00:00", "a", "甲"), ("01:20:00", "b", "乙")))
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    cues, _ = exporters.times_for(pairs, None)
    assert cues[0][1] - cues[0][0] <= exporters.MAX_CUE_S + 1e-9


def test_times_for_handles_unparsable_clock():
    main.TRANSCRIPTS.mkdir(parents=True, exist_ok=True)
    p = main.TRANSCRIPTS / NOTE
    p.write_text("- `bad-ts` EN one\n  - `#1` 中文一\n", encoding="utf-8")
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    cues, _ = exporters.times_for(pairs, None)
    assert len(cues) == 1 and cues[0][1] >= exporters.MIN_CUE_S


# —— 时间轴：precise（sidecar） ——

def test_times_path_and_roundtrip():
    p = write_note()
    segs = [{"seg": 1, "start": 0.4, "end": 2.9},
            {"seg": 2, "start": 3.0, "end": 5.5}]
    exporters.write_times(p, t0=123.0, segments=segs)
    assert exporters.times_path(p).name == "lecture_20260101_0900.times.json"
    got = exporters.read_times(p)
    assert got["t0"] == 123.0 and got["segments"] == segs
    assert not exporters.times_path(p).with_name(
        exporters.times_path(p).name + ".tmp").exists()     # 原子写不留 .tmp


def test_read_times_tolerates_missing_and_corrupt():
    p = write_note()
    assert exporters.read_times(p) is None                   # 未写过
    tp = exporters.times_path(p)
    tp.write_text("{ not json", encoding="utf-8")
    assert exporters.read_times(p) is None
    tp.write_text(json.dumps({"version": 999, "segments": []}), encoding="utf-8")
    assert exporters.read_times(p) is None                   # 版本不符
    tp.write_text(json.dumps({"version": 1, "segments": "nope"}), encoding="utf-8")
    assert exporters.read_times(p) is None


def test_times_for_precise_uses_sidecar():
    p = write_note()
    exporters.write_times(p, 0.0, [{"seg": 1, "start": 0.4, "end": 2.9},
                                   {"seg": 2, "start": 3.0, "end": 5.5}])
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    cues, source = exporters.times_for(pairs, exporters.read_times(p))
    assert source == "precise"
    assert cues[0] == (0.4, 2.9) and cues[1] == (3.0, 5.5)


def test_times_for_precise_fills_missing_seg_from_derived():
    """sidecar 缺某段（如中途崩溃）→ 该段退回墙钟差分，不整体作废。

    段值须与 derived 基线时间上自洽，否则会被单调钳制改写（那是正确行为：
    见 test_times_for_derived_clamps_same_second_segments）。"""
    p = write_note(pairs=(("01:00:00", "a", "甲"), ("01:00:05", "b", "乙"),
                          ("01:00:10", "c", "丙")))
    exporters.write_times(p, 0.0, [{"seg": 3, "start": 10.2, "end": 12.8}])
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    cues, source = exporters.times_for(pairs, exporters.read_times(p))
    assert source == "precise"
    assert cues[2] == (10.2, 12.8)                            # 有 sidecar 的那段用精确值
    assert cues[0] == (0.0, 5.0)                              # 缺的段退回墙钟差分
    assert cues[1] == (5.0, 10.0)


def test_times_for_empty_pairs():
    assert exporters.times_for([], None) == ([], "derived")


# —— SRT 产物 ——

def test_build_srt_bilingual_block_shape():
    p = write_note()
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    srt, source = exporters.build_srt(pairs)
    assert source == "derived"
    blocks = srt.strip().split("\n\n")
    assert len(blocks) == 2
    first = blocks[0].splitlines()
    assert first[0] == "1"
    assert " --> " in first[1]
    assert first[2] == "Hello world."        # 上英
    assert first[3] == "你好世界。"            # 下中
    assert blocks[1].startswith("2\n")


def test_build_srt_english_only_single_line():
    p = write_note(pairs=(("01:00:00", "Only english.", None),))
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    srt, _ = exporters.build_srt(pairs)
    assert srt.splitlines()[2] == "Only english."
    assert len(srt.strip().split("\n\n")) == 1


def test_build_srt_empty_is_empty_string():
    srt, _ = exporters.build_srt([])
    assert srt == ""


# —— PDF ——

def _pdf_ok(data: bytes) -> bool:
    return data[:5] == b"%PDF-" and b"%%EOF" in data[-2048:]


def test_build_pdf_real_bytes():
    p = write_note()
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    data = exporters.build_pdf(pairs, {"title": "双语课堂笔记 20260101_0900",
                                       "subtitle": "课程：计算机科学导论 · 2 段"})
    assert isinstance(data, bytes) and len(data) > 1000
    assert _pdf_ok(data)
    assert b"/FontFile2" in data              # 字体真嵌入（中文才可能显示）
    assert b"/Type0" in data                  # CID 字体（CJK 必需）


def test_build_pdf_missing_font_raises_visible_error(monkeypatch):
    """字体缺失必须抛 **RuntimeError 且带可操作指引**，而不是静默出豆腐块 PDF。

    评审1 修前：旧实现只查 FONT_BOLD，且抛的是 fpdf2 原生 FileNotFoundError
    （既不是 RuntimeError 也没有指引），而测试只断言 pytest.raises(Exception)，
    所以这条契约无人守。现在把类型与指引文案一起钉住。"""
    monkeypatch.setattr(exporters, "FONT_REGULAR", "C:/definitely/not/a/font.ttf")
    p = write_note()
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    with pytest.raises(RuntimeError) as ei:
        exporters.build_pdf(pairs, {"title": "t", "subtitle": ""})
    msg = str(ei.value)
    assert "字体缺失" in msg and "微软雅黑" in msg
    assert exporters.FONT_REGULAR in msg          # 把实际路径告诉用户


def test_build_pdf_missing_bold_font_falls_back(monkeypatch, tmp_path):
    """粗体是可选的：缺失只回落常规字面，不得让导出失败。"""
    monkeypatch.setattr(exporters, "FONT_BOLD", "C:/definitely/not/a/bold.ttf")
    p = write_note()
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    assert _pdf_ok(exporters.build_pdf(pairs, {"title": "t", "subtitle": ""}))


def test_build_pdf_handles_empty_and_english_only():
    data = exporters.build_pdf([], {"title": "空", "subtitle": ""})
    assert _pdf_ok(data)
    p = write_note(pairs=(("01:00:00", "English only.", None),))
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    assert _pdf_ok(exporters.build_pdf(pairs, {"title": "t", "subtitle": ""}))


# —— DOCX ——

def test_build_docx_real_bytes_and_eastasia_font():
    p = write_note()
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    data = exporters.build_docx(pairs, {"title": "双语课堂笔记", "subtitle": "课程：甲"})
    assert data[:2] == b"PK"                  # docx = zip
    # 关键：中文必须带 w:eastAsia，否则 Word 回退字体。
    # docx 是 zip，document.xml 被 deflate 压缩——必须解压后再查，直接搜裸字节查不到。
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml")
    assert b"w:eastAsia" in xml
    assert exporters.DOCX_EA_FONT.encode() in xml


def test_build_docx_roundtrip_content():
    import docx as docx_mod
    p = write_note()
    pairs = notes_mod.parse(p.read_text(encoding="utf-8"))["pairs"]
    data = exporters.build_docx(pairs, {"title": "标题测试", "subtitle": "副标题"})
    d = docx_mod.Document(io.BytesIO(data))
    text = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Hello world." in text and "你好世界。" in text
    assert "\n".join(x.text for x in d.paragraphs).find("标题测试") >= 0


def test_build_docx_empty():
    assert exporters.build_docx([], {"title": "空的", "subtitle": ""})[:2] == b"PK"


# —— asr.transcribe 三形态 ——

class _Seg:
    def __init__(self, start, end, text):
        self.start, self.end, self.text = start, end, text


def _make_asr(segments):
    from asr import WhisperASR
    w = object.__new__(WhisperASR)            # 绕过 __init__：不加载真模型
    w.language = "en"
    w.detected_lang = None
    w.backend = "ctranslate2"                 # Mac 适配：transcribe 按 backend 分发

    class M:
        def transcribe(self, audio, **kw):
            return iter(segments), SimpleNamespace(language="en")

    w.model = M()
    return w


def test_transcribe_with_times_returns_segments():
    w = _make_asr([_Seg(0.0, 1.2, " hello"), _Seg(1.2, 2.0, " world")])
    text, segs = w.transcribe(np.zeros(16000, dtype=np.float32), with_times=True)
    assert text == "hello world"
    assert segs == [{"start": 0.0, "end": 1.2, "text": "hello"},
                    {"start": 1.2, "end": 2.0, "text": "world"}]


def test_transcribe_default_unchanged_and_word_timestamps_unchanged():
    """向后兼容：默认仍返回 str；word_timestamps 仍返回原始 segments 列表。"""
    w = _make_asr([_Seg(0.0, 1.0, " hi")])
    assert w.transcribe(np.zeros(16000, dtype=np.float32)) == "hi"
    w2 = _make_asr([_Seg(0.0, 1.0, " hi")])
    got = w2.transcribe(np.zeros(16000, dtype=np.float32), word_timestamps=True)
    assert isinstance(got, list) and not isinstance(got, tuple)


def test_transcribe_with_times_skips_blank_segments():
    w = _make_asr([_Seg(0.0, 0.5, "  "), _Seg(0.5, 1.0, " real")])
    text, segs = w.transcribe(np.zeros(16000, dtype=np.float32), with_times=True)
    assert text == "real" and len(segs) == 1


# —— 端点 ——

async def test_export_md_srt_pdf_docx_end_to_end():
    write_note()
    async with asgi_client() as c:
        for fmt, magic in (("md", b"#"), ("srt", b"1"), ("pdf", b"%PDF-"), ("docx", b"PK")):
            r = await c.get("/api/notes/export",
                            params={"name": NOTE, "fmt": fmt})
            assert r.status_code == 200, f"{fmt}: {r.text[:200]}"
            assert r.content[:len(magic)] == magic, f"{fmt} 产物头不对"
            assert "attachment" in r.headers.get("content-disposition", "")
        r = await c.get("/api/notes/export", params={"name": NOTE, "fmt": "srt"})
        assert "Hello world." in r.text and "你好世界。" in r.text


async def test_export_defaults_to_md():
    write_note()
    async with asgi_client() as c:
        r = await c.get("/api/notes/export", params={"name": NOTE})
        assert r.status_code == 200 and "双语课堂笔记" in r.text


async def test_export_invalid_fmt_400_with_message():
    write_note()
    async with asgi_client() as c:
        r = await c.get("/api/notes/export", params={"name": NOTE, "fmt": "exe"})
        assert r.status_code == 400
        assert "fmt" in r.json()["error"] or "格式" in r.json()["error"]


async def test_export_missing_note_404():
    async with asgi_client() as c:
        r = await c.get("/api/notes/export",
                        params={"name": "lecture_20990101_0000.md", "fmt": "srt"})
        assert r.status_code == 404


async def test_export_traversal_rejected():
    async with asgi_client() as c:
        for bad in ("../config.json", "C:/Windows/win.ini", "config.json"):
            r = await c.get("/api/notes/export", params={"name": bad, "fmt": "pdf"})
            assert r.status_code == 404, f"应拒绝 {bad}"


async def test_export_failure_returns_visible_error(monkeypatch):
    """生成失败要给**用户可见**的错误体（JSON error），不是裸 500 空响应。"""
    write_note()

    def boom(*a, **kw):
        raise RuntimeError("font exploded")

    monkeypatch.setattr(exporters, "build_pdf", boom)
    monkeypatch.setattr(main, "exporters", exporters)
    async with asgi_client() as c:
        r = await c.get("/api/notes/export", params={"name": NOTE, "fmt": "pdf"})
    assert r.status_code == 500
    body = r.json()
    assert body["ok"] is False and "font exploded" in body["error"]
    assert "pdf" in body["error"]


# —— p2-3 时间戳贯通：解包 / 调度器 meta / Hub sidecar ——

def test_unpack_transcribe_both_shapes():
    assert main.unpack_transcribe(("text", [{"start": 0, "end": 1}])) == \
        ("text", [{"start": 0, "end": 1}])
    assert main.unpack_transcribe("plain") == ("plain", [])      # 替身/旧实现
    assert main.unpack_transcribe(None) == ("", [])


class _TimeASR:
    """返回 (text, segs) 的替身，模拟 asr.transcribe(with_times=True)。"""
    name, device = "t", "cpu"

    def transcribe(self, audio, initial_prompt=None, **kw):
        if kw.get("with_times"):
            return "hello world", [{"start": 0.5, "end": 1.5, "text": "hello"},
                                   {"start": 1.5, "end": 2.5, "text": "world"}]
        return "hello world"


class _PlainASR:
    """只回 str 的替身 → 必须走空 seg_times 且不抛。"""
    name, device = "p", "cpu"

    def transcribe(self, audio, initial_prompt=None, **kw):
        return "plain text"


def _run_scheduler_final(asr):
    import scheduler as sched_mod
    rec: dict = {}

    def on_final(en, meta):
        rec["en"], rec["meta"] = en, dict(meta)

    s = sched_mod.DecodeScheduler(asr_provider=lambda: asr,
                                  on_final=on_final, emit_partial=lambda p: None)
    s.enqueue_final(asr, np.zeros(16000, dtype=np.float32), ts_audio=10.0)
    for _ in range(50):
        if not s.step(timeout=0.0):
            break
    return rec


def test_scheduler_meta_carries_seg_times():
    rec = _run_scheduler_final(_TimeASR())
    assert rec["en"] == "hello world"
    assert rec["meta"]["seg_times"] == [
        {"start": 0.5, "end": 1.5, "text": "hello"},
        {"start": 1.5, "end": 2.5, "text": "world"}]
    assert rec["meta"]["audio_s"] == pytest.approx(1.0)


def test_scheduler_meta_seg_times_empty_for_str_only_asr():
    """向后兼容：只回 str 的 ASR 不得让 meta 缺键或抛异常。"""
    rec = _run_scheduler_final(_PlainASR())
    assert rec["en"] == "plain text"
    assert rec["meta"]["seg_times"] == []


def test_hub_record_time_offsets_and_sidecar():
    """第 2 个参数是**窗口起点**（评审1 修后的口径），不是 capture 的发射时刻。"""
    hub = main.Hub()
    p = hub.start_transcript("甲")
    hub.record_time(1, 1000.0, [{"start": 0.5, "end": 2.5}], 3.0)
    hub.record_time(2, 1004.0, [{"start": 0.2, "end": 1.8}], 2.0)
    assert hub.t0_audio == 1000.0                     # 原点 = 首段窗口起点
    assert hub.times == [{"seg": 1, "start": 0.5, "end": 2.5},
                         {"seg": 2, "start": 4.2, "end": 5.8}]
    got = exporters.read_times(p)                     # 同步落盘
    assert got["t0"] == 1000.0 and got["segments"] == hub.times


def test_record_time_anchor_is_window_start_not_emit_time():
    """回归（评审1 核心）：锚点必须是**窗口起点**，不是 capture 的发射时刻。

    capture 的 `q_out.put((time.time(), audio))` 在 `_emit` 末尾执行——那一刻整段
    缓冲（0.4s preroll + 语音 + ~0.7s 尾随静音）已组装完，故时间戳 ≈ 窗口**末尾**。
    旧实现拿它当起点，第 i 段起点误差恰为 `d_i − d_1`（d = 窗口时长），SRT 每条 cue
    整体后移、相邻段还被 _clamp_cues 压成最短 cue。

    本用例按真值回放：给定各段的「窗口起点 + 窗内 whisper 偏移」，造出 capture 会
    投递的发射时刻，再用修好的换算喂进 record_time，断言结果与真值逐条相等。"""
    # (窗口起点, 窗内 start, 窗内 end, 窗口时长)
    truth = [(0.0, 0.4, 5.3, 6.0),
             (7.0, 0.4, 1.6, 2.0),
             (10.0, 0.4, 6.4, 7.0),
             (18.0, 0.4, 2.8, 3.0)]
    hub = main.Hub()
    hub.start_transcript()
    for i, (win_start, s, e, dur) in enumerate(truth, start=1):
        emit_ts = win_start + dur                     # capture 实际投递的就是这个
        got_win_start = emit_ts - dur                 # 调用方换算回窗口起点
        hub.record_time(i, got_win_start,
                        [{"start": s, "end": e, "text": "x"}], dur)
    expect = [{"seg": i, "start": round(ws + s, 3), "end": round(ws + e, 3)}
              for i, (ws, s, e, _) in enumerate(truth, start=1)]
    assert hub.times == expect, "cue 起点没有锚到窗口起点——每段都整体后移了"


def test_old_emit_time_anchor_would_misplace_cues():
    """对照组：若把发射时刻当起点（旧实现），误差正好是 `d_i − d_1`。

    这条不测产品代码，而是把「为什么必须换锚点」钉成可读的数字，防止有人把换算改回去。"""
    truth = [(0.0, 6.0), (7.0, 2.0), (10.0, 7.0), (18.0, 3.0)]   # (窗口起点, 时长)
    emit = [ws + d for ws, d in truth]
    d1 = truth[0][1]
    wrong = [(e - emit[0]) for e in emit]                          # 旧口径的 off
    errors = [w - ws for w, (ws, _) in zip(wrong, truth)]
    assert errors == [0.0, -4.0, 1.0, -3.0]
    assert errors == [d - d1 for _, d in truth]                    # 即 d_i − d_1


def test_hub_record_time_without_seg_times_uses_audio_window():
    hub = main.Hub()
    p = hub.start_transcript()
    hub.record_time(1, 500.0, [], 4.0)
    assert hub.times == [{"seg": 1, "start": 0.0, "end": 4.0}]


def test_hub_record_time_zero_length_and_clock_skew():
    """时钟回拨 + 零长段：不得出负值/零长区间。

    跨段单调性**不在这里**保证（此处如实记原始值，导出侧 _clamp_cues 统一收口）。"""
    hub = main.Hub()
    hub.start_transcript()
    hub.record_time(1, 900.0, None, 3.0)
    hub.record_time(2, 800.0, [{"start": 0.0, "end": 0.0}], 2.0)   # 回拨 + 零长
    for t in hub.times:
        assert t["start"] >= 0.0
        assert t["end"] > t["start"]


def test_hub_record_time_noop_without_audio_ts():
    hub = main.Hub()
    hub.start_transcript()
    hub.record_time(1, None, [], 3.0)
    assert hub.times == [] and hub.t0_audio is None


def test_start_transcript_resets_times():
    hub = main.Hub()
    hub.start_transcript()
    hub.record_time(1, 100.0, [], 2.0)
    assert hub.times
    hub.start_transcript()                            # 换笔记 = 时间轴重新起算
    assert hub.times == [] and hub.t0_audio is None


def test_hub_record_time_survives_unwritable_sidecar(monkeypatch):
    """sidecar 写不进去也只告警：绝不能因为时间戳写盘失败打断字幕/笔记主链路。"""
    hub = main.Hub()
    hub.start_transcript()

    def boom(*a, **kw):
        raise OSError("disk full")

    monkeypatch.setattr(exporters, "write_times", boom)
    hub.record_time(1, 100.0, [], 2.0)                # 不抛
    assert hub.times and len(hub.times) == 1


# —— 端到端：asr_loop → hub.record_time → sidecar ——
# 这是时间戳链路里唯一没被其他用例覆盖的一环（scheduler/Hub 已各自单测），
# 故用打桩的重型依赖把**真实主循环**跑起来，喂一段音频走完全程。

class _LoopASR:
    def transcribe(self, audio, initial_prompt=None, **kw):
        if kw.get("with_times"):
            return "hello world", [{"start": 0.25, "end": 1.75, "text": "hello world"}]
        return "hello world"


class _LoopSegmenter:
    # **kw：p2-5 起 apply_segmenter 还会传断句参数（真 MicSegmenter 收这三个）
    def __init__(self, q_out, device=None, vad="silero", stream_q=None, **kw):
        self.error, self.thread, self.last_feed_ts = None, None, 0.0

    def start(self): pass

    def stop(self): pass

    def snapshot(self): return {}


def _loop_args():
    return SimpleNamespace(model=None, gpu=False, compute_type=None, vad=None,
                           streaming=False, lang=None, api_key="", glossary=[],
                           source=None, device=None)


async def test_asr_loop_records_absolute_times_to_sidecar(monkeypatch):
    """真实主循环跑一段音频：t0 = 首段采集墙钟，段区间 = t0 + whisper 段级时间，
    并且 sidecar 已落盘（导出侧 SRT 的 precise 档就吃这个文件）。"""
    saved = {k: main.runtime.get(k) for k in
             ("asr", "segmenter", "seg_q", "scheduler", "translator", "prompt")}
    saved_paused = main.status["paused"]
    main.runtime["config"].clear()
    main.runtime["config"].update(main.DEFAULT_CONFIG)
    main.runtime["config"]["streaming"] = False     # 走整句路径（scheduler is None）
    for k in ("asr", "scheduler", "translator"):
        main.runtime[k] = None
    main.status["paused"] = False
    hub = main.Hub()
    monkeypatch.setattr(main, "hub", hub)
    p = hub.start_transcript("甲")
    monkeypatch.setattr(main, "MicSegmenter", _LoopSegmenter)
    monkeypatch.setattr(main, "ARGS", SimpleNamespace(device=None))
    monkeypatch.setattr(main, "load_api_key", lambda *a, **k: "")

    def _build(overrides=None):
        main.runtime["asr"] = _LoopASR()
        main.status["asr"] = "ready"
    monkeypatch.setattr(main, "build_asr", _build)

    async def _noop():
        return
    monkeypatch.setattr(main, "status_loop", _noop)

    task = asyncio.create_task(main.asr_loop(_loop_args()))
    try:
        for _ in range(300):
            if main.runtime["seg_q"] is not None and main.runtime["asr"] is not None:
                break
            await asyncio.sleep(0.01)
        main.runtime["seg_q"].put((5000.0, np.zeros(16000 * 2, dtype=np.float32)))
        for _ in range(400):
            if hub.times:
                break
            await asyncio.sleep(0.01)
    finally:
        # 收尾必须照 lifespan 那套来：asr_loop 会阻塞在 to_thread(seg_q.get) 上，
        # 只 cancel 的话执行器线程会永远卡在 get() —— ThreadPoolExecutor 的线程在
        # 解释器退出时被 join，表现为「测试全绿但进程挂死」（实测踩过：单跑本文件
        # 最后一项之后不退出）。故：先置 paused 让循环改走暂停分支（不再 get），
        # 再投哨兵解除当前阻塞的 get，最后 cancel。
        main.status["paused"] = True
        task.cancel()
        if main.runtime["seg_q"] is not None:
            main.runtime["seg_q"].put(None)
        try:
            await task
        except asyncio.CancelledError:
            pass
        for k, v in saved.items():
            main.runtime[k] = v
        main.status["paused"] = saved_paused

    assert hub.times, "主循环没有记录段级时间戳——asr_loop → record_time 未接通"
    # 喂进去的发射时刻是 5000.0，音频 2s → 窗口起点 4998.0（评审1 修后的口径）
    assert hub.t0_audio == 4998.0
    assert hub.times[0] == {"seg": 1, "start": 0.25, "end": 1.75}
    # 主循环启动时**自己**开一篇笔记（本用例预建的同分钟笔记会被让路成 -2，见
    # Hub.start_transcript 的重名避让），sidecar 落盘在它自己那篇旁边——故断言读
    # hub.transcript_path，而不是用例开头预建的 p。
    got = exporters.read_times(hub.transcript_path)
    assert got and got["segments"] == hub.times, "sidecar 未落盘"
    assert hub.transcript_path != p, "同分钟重名未让路（会把预建笔记整篇截掉）"


async def test_export_srt_prefers_sidecar_when_present():
    p = write_note()
    exporters.write_times(p, 0.0, [{"seg": 1, "start": 0.4, "end": 2.9},
                                   {"seg": 2, "start": 3.0, "end": 5.5}])
    async with asgi_client() as c:
        r = await c.get("/api/notes/export", params={"name": NOTE, "fmt": "srt"})
    assert "00:00:00,400 --> 00:00:02,900" in r.text
    assert "00:00:03,000 --> 00:00:05,500" in r.text
