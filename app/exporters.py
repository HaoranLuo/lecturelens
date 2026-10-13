# -*- coding: utf-8 -*-
"""exporters.py — p2-3 导出中心：md / SRT / PDF / Word 四种产物。

设计要点
--------
**SRT 时间来源（两档，按可信度降级）**
1. `precise`：<笔记>.times.json sidecar —— 每段音频采集墙钟 + faster-whisper 段级
   start/end 换算出的**绝对偏移秒**（main.Hub.record_time 落盘）。这是真实语音边界。
2. `derived`：笔记行自带的 `HH:MM:SS` 墙钟差分（下一段开始即本段结束）。所有历史
   笔记都有这个（p2-2 起解析器已给出），故**任何一篇笔记都能导出 SRT**，只是精度为秒级。
两条路都强制单调递增且不短于 MIN_CUE_S，避免秒级分辨率的同秒段退化成 0 时长 cue。

**打印版配色 ≠ 应用内令牌（对 docs/design-spec.md §二 的有意调整）**
规范 §二 的 `--gold: #FFD97A` 是深色界面上的品牌色；白纸上它的对比度约
**1.3:1**，远低于规范自己在同节写明的「对比度均 ≥ 4.5:1」。PDF/Word 是白纸打印
件，故中文译文改用同一金色族的深色变体 `#8A6A14`（对白底实测约 **5.0:1**），
英文用近黑、时间戳用灰——保持「英文原文 + 中文金色译文」的语义分层不变。
这是硬规矩⑨意义上的按事实调整，非无视规范。

**依赖**：fpdf2（模块名 fpdf）与 python-docx（模块名 docx）在 requirements.txt
的 P2 段；PDF 中文必须注册系统字体 msyh.ttc（见 main 的字体常量与 p2-deps 结论）。
"""
import io
import json
import logging
from pathlib import Path

log = logging.getLogger("exporters")

#: 系统字体（PDF 中文必需；fpdf2 会子集化嵌入，勿整字库嵌入）
FONT_REGULAR = "C:/Windows/Fonts/msyh.ttc"
FONT_BOLD = "C:/Windows/Fonts/msyhbd.ttc"

#: 打印版配色（见模块 docstring 的对比度说明）
INK = (26, 32, 40)          # 英文原文/标题
MUTED = (110, 122, 134)     # 时间戳、元信息
ZH_INK = (138, 106, 20)     # 中文译文（金色族深色变体，白底 ≈5.0:1）
RULE = (208, 214, 221)      # 分隔线

MIN_CUE_S = 0.6             # SRT 单条最小时长（秒）
DEFAULT_CUE_S = 3.0         # 末段无后继可参考时的默认时长
MAX_CUE_S = 30.0            # 单条上限：防止长静默把一条拖成整节课

TIMES_SUFFIX = ".times.json"
TIMES_VERSION = 1


# —— 时间 sidecar 读写（写入方 = main.Hub，读取方 = 导出） ——

def times_path(note_path: Path) -> Path:
    """某篇笔记对应的时间 sidecar 路径（与笔记同目录同名 + .times.json）。"""
    p = Path(note_path)
    return p.with_name(p.stem + TIMES_SUFFIX)


def write_times(note_path: Path, t0: float | None, segments: list) -> None:
    """原子写时间 sidecar：{"version", "t0", "segments":[{"seg","start","end"}]}。

    t0 = 本场第一段音频的采集墙钟（绝对偏移的零点）。写失败只告警不回抛——
    导出是锦上添花，绝不能因为写时间戳失败而打断字幕/笔记主链路。"""
    p = times_path(note_path)
    try:
        tmp = p.with_name(p.name + ".tmp")
        tmp.write_text(json.dumps(
            {"version": TIMES_VERSION, "t0": t0, "segments": segments or []},
            ensure_ascii=False), encoding="utf-8")
        tmp.replace(p)
    except Exception as e:
        log.warning(f"[export] ⚠ 时间戳 sidecar 写入失败（不影响导出降级路径）: {e}")


def read_times(note_path: Path) -> dict | None:
    """读时间 sidecar；缺失/损坏/版本不符 → None（调用方回退墙钟差分）。"""
    p = times_path(note_path)
    try:
        if not p.is_file():
            return None
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None
    if not isinstance(data, dict) or data.get("version") != TIMES_VERSION:
        return None
    segs = data.get("segments")
    if not isinstance(segs, list):
        return None
    return {"t0": data.get("t0"), "segments": segs}


# —— SRT 时间轴 ——

def srt_timestamp(seconds: float) -> str:
    """秒 → SRT 时间戳 `HH:MM:SS,mmm`。负数按 0 处理（时钟回拨等异常输入）。"""
    s = max(0.0, float(seconds))
    ms = int(round(s * 1000))
    h, rem = divmod(ms, 3600_000)
    m, rem = divmod(rem, 60_000)
    sec, msec = divmod(rem, 1000)
    return f"{h:02d}:{m:02d}:{sec:02d},{msec:03d}"


def clock_seconds(ts: str) -> float | None:
    """笔记行的 `HH:MM:SS` → 当日秒数；非法 → None。"""
    if not isinstance(ts, str):
        return None
    parts = ts.strip().split(":")
    if len(parts) != 3:
        return None
    try:
        h, m, s = (int(x) for x in parts)
    except ValueError:
        return None
    if not (0 <= h < 24 and 0 <= m < 60 and 0 <= s < 60):
        return None
    return h * 3600 + m * 60 + s


def _clamp_cues(cues: list) -> list:
    """把 (start,end) 序列修成合法 cue：单调递增、不短于 MIN_CUE_S、不长于 MAX_CUE_S。

    墙钟差分在秒级分辨率下会让相邻段同秒（差 0），裸用会产出 0 时长 cue
    （多数播放器直接不显示）——故必须修。"""
    out: list = []
    prev_end = 0.0
    for i, (st, en) in enumerate(cues):
        s = max(float(st), prev_end)
        e = max(float(en), s + MIN_CUE_S)
        if e - s > MAX_CUE_S:
            e = s + MAX_CUE_S
        out.append((s, e))
        prev_end = e
    return out


def times_for(pairs: list, sidecar: dict | None) -> tuple[list, str]:
    """算出每段的 (start, end) 秒 + 来源标记（"precise" / "derived"）。

    precise：sidecar 按 seg 对齐（缺哪段就退回该段用 derived 的值兜底）；
    derived：由笔记行墙钟差分。首段起点恒为 0（导出的时间轴以本场第一段为原点）。"""
    n = len(pairs)
    if not n:
        return [], "derived"

    # derived 基线（总是先算，兼作 precise 的缺段填充）
    secs = [clock_seconds(p.get("ts", "")) for p in pairs]
    base: list[tuple[float, float]] = []
    origin = next((s for s in secs if s is not None), None)
    for i in range(n):
        s = secs[i]
        st = 0.0 if (s is None or origin is None) else max(0.0, s - origin)
        nxt = next((secs[j] for j in range(i + 1, n) if secs[j] is not None), None)
        en = (nxt - origin) if (nxt is not None and origin is not None) else st + DEFAULT_CUE_S
        base.append((st, en))
    base = _clamp_cues(base)

    if not sidecar:
        return base, "derived"
    by_seg: dict = {}
    for item in sidecar.get("segments") or []:
        if not isinstance(item, dict):
            continue
        try:
            by_seg[int(item["seg"])] = (float(item["start"]), float(item["end"]))
        except (KeyError, TypeError, ValueError):
            continue
    if not by_seg:
        return base, "derived"
    merged = []
    for i, p in enumerate(pairs):
        seg = p.get("seg", i + 1)
        merged.append(by_seg.get(int(seg), base[i]))
    return _clamp_cues(merged), "precise"


def build_srt(pairs: list, sidecar: dict | None = None,
              bilingual: bool = True) -> tuple[str, str]:
    """生成 SRT 文本 → (srt, source)。bilingual=True 时同一条 cue 两行（上英下中）。

    纯英文笔记（无 ZH）自动退化为单行，不留空行。"""
    cues, source = times_for(pairs, sidecar)
    blocks: list[str] = []
    idx = 0
    for pair, (t_start, t_end) in zip(pairs, cues):
        text_en = (pair.get("en") or "").strip()
        text_zh = (pair.get("zh") or "").strip()
        if not text_en and not text_zh:
            continue
        idx += 1
        lines = [text_en]
        if bilingual and text_zh:
            lines.append(text_zh)
        text = "\n".join(x for x in lines if x)
        blocks.append(f"{idx}\n{srt_timestamp(t_start)} --> "
                      f"{srt_timestamp(t_end)}\n{text}")
    return "\n\n".join(blocks) + ("\n" if blocks else ""), source


# —— PDF（fpdf2 + 系统字体） ——

def _register_pdf_fonts(pdf) -> tuple[str, str]:
    """注册微软雅黑（含粗体）。返回 (常规字体名, 粗体字体名)。

    字体缺失时抛 **RuntimeError 带明确指引**——导出失败必须让用户看见可操作的原因，
    而不是静默出一个豆腐块 PDF。

    评审1 修：旧实现只检查 FONT_BOLD、不检查 FONT_REGULAR，且抛的是 fpdf2 原生的
    `FileNotFoundError: TTF Font file not found: ...`——既不是 RuntimeError、也没有
    任何指引（中文界面里就露出一句英文异常），与本函数 docstring 的承诺不符。现按
    docstring 兑现：常规字面缺失 → RuntimeError + 可操作指引；粗体缺失 → 回落常规。"""
    reg, bold = "msyh", "msyhbd"
    if not Path(FONT_REGULAR).is_file():
        raise RuntimeError(
            f"PDF 中文字体缺失：找不到 {FONT_REGULAR}。"
            f"请确认运行环境是 Windows 且已安装微软雅黑（Microsoft YaHei），"
            f"或把 exporters.FONT_REGULAR 指向本机可用的中文字体（.ttf/.ttc）。")
    try:
        pdf.add_font(reg, "", FONT_REGULAR)
    except Exception as e:
        raise RuntimeError(
            f"注册中文字体失败：{FONT_REGULAR}（{type(e).__name__}: {e}）。"
            f"请确认该文件是有效的 TTF/TTC 字体。") from e
    if Path(FONT_BOLD).is_file():
        pdf.add_font(bold, "", FONT_BOLD)
    else:
        bold = reg            # 粗体可选：缺失就回落常规字面，不阻断导出
    return reg, bold


def build_pdf(pairs: list, meta: dict, sidecar: dict | None = None) -> bytes:
    """双语对照 PDF：标题 + 元信息 + 每段「时间 | 英文 | 中文」三栏。

    版式对齐 docs/design-spec.md §七-2 的双语对照观感（时间戳小字、英文在上、
    中文金色），并按 A4 打印重排（配色见模块 docstring）。"""
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.set_margins(18, 16, 18)
    pdf.add_page()
    reg, bold = _register_pdf_fonts(pdf)

    avail = pdf.w - pdf.l_margin - pdf.r_margin      # A4 竖版约 174mm
    w_time, gap = 18.0, 4.0
    w_col = (avail - w_time - 2 * gap) / 2

    # 标题 + 元信息
    pdf.set_font(bold, size=16)
    pdf.set_text_color(*INK)
    pdf.multi_cell(avail, 9, meta.get("title") or "双语课堂笔记")
    pdf.set_font(reg, size=9)
    pdf.set_text_color(*MUTED)
    pdf.multi_cell(avail, 5.5, meta.get("subtitle") or "")
    pdf.ln(1.5)
    y0 = pdf.get_y()
    pdf.set_draw_color(*RULE)
    pdf.line(pdf.l_margin, y0, pdf.l_margin + avail, y0)
    pdf.set_y(y0 + 2.5)

    def col_height(text: str, w: float, size: float) -> float:
        # 必须连字族一起设：只改 size 会沿用上一处的字族（标题用的粗体），
        # 量出的高度会有偏差 → 三栏对不齐。
        pdf.set_font(reg, size=size)
        return pdf.multi_cell(w, 5.2, text, dry_run=True, output="HEIGHT")

    bottom = pdf.h - pdf.b_margin
    for p in pairs:
        en = (p.get("en") or "").strip()
        zh = (p.get("zh") or "").strip()
        if not en and not zh:
            continue
        h = max(col_height(en or " ", w_col, 10),
                col_height(zh or " " , w_col, 10), 5.2)
        if pdf.get_y() + h > bottom:                 # 手动分页（auto_page_break 已关）
            pdf.add_page()
            pdf.set_font_size(10)
        x, y = pdf.l_margin, pdf.get_y()
        pdf.set_font(reg, size=8)
        pdf.set_text_color(*MUTED)
        pdf.set_xy(x, y)
        pdf.multi_cell(w_time, 5.2, p.get("ts") or "")
        pdf.set_font(reg, size=10)
        pdf.set_text_color(*INK)
        pdf.set_xy(x + w_time + gap, y)
        pdf.multi_cell(w_col, 5.2, en)
        pdf.set_text_color(*ZH_INK)
        pdf.set_xy(x + w_time + gap + w_col + gap, y)
        pdf.multi_cell(w_col, 5.2, zh)
        pdf.set_y(y + h)

    out = pdf.output()
    return bytes(out)


# —— Word（python-docx） ——

DOCX_EA_FONT = "Microsoft YaHei"


def _docx_run(paragraph, text: str, *, size=None, color=None, bold=False):
    """加一个 run 并**显式设定东亚字体**。

    python-docx 只设 font.name 时写的是 w:ascii/w:hAnsi，Word 对中日韩文字会回退
    到默认宋体（在部分环境甚至回退成非中文字体）——必须同时写 w:eastAsia。
    （p2-deps 结论里我把这条列为「未核实」，本项落地时按此实现。）"""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run = paragraph.add_run(text)
    run.font.name = DOCX_EA_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), DOCX_EA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    return run


def build_docx(pairs: list, meta: dict, sidecar: dict | None = None) -> bytes:
    """双语对照 .docx：与 PDF 同构（标题 + 元信息 + 三栏表格）。"""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_heading(meta.get("title") or "双语课堂笔记", level=1)
    sub = doc.add_paragraph()
    _docx_run(sub, meta.get("subtitle") or "", size=9, color=MUTED)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Cm(2.2), Cm(7.4), Cm(7.4))
    hdr = table.rows[0].cells
    for cell, label, w in zip(hdr, ("时间", "英文原文", "中文译文"), widths):
        cell.width = w
        _docx_run(cell.paragraphs[0], label, size=9, color=MUTED, bold=True)

    for p in pairs:
        en = (p.get("en") or "").strip()
        zh = (p.get("zh") or "").strip()
        if not en and not zh:
            continue
        cells = table.add_row().cells
        for cell, w in zip(cells, widths):
            cell.width = w
        _docx_run(cells[0].paragraphs[0], p.get("ts") or "", size=8, color=MUTED)
        _docx_run(cells[1].paragraphs[0], en, size=10, color=INK)
        _docx_run(cells[2].paragraphs[0], zh, size=10, color=ZH_INK)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
